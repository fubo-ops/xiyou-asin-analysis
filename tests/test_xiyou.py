from datetime import date
import re
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from xiyou_skill.models import CaseConfig, RawArtifact
from xiyou_skill.xiyou import normalize_xiyou_xlsx


def test_normalize_xiyou_maps_known_fields_and_cell_lineage(tmp_path):
    path = tmp_path / "TEST_FIXTURE_not_real_data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "多ASIN总览"
    ws.append(["ASIN", "7天流量得分", "预估30天订单", "自然关键词数"])
    ws.append(["B0SYNTH001", 88.5, 42, 120])
    ws.append(["B0SYNTH002", 120.0, 77, 210])
    wb.save(path)
    case = CaseConfig("case-1", "US", "B0SYNTH001", ("B0SYNTH002",), date(2026, 7, 12), date(2026, 8, 10))
    artifact = RawArtifact("X1", "XIYOU", "L3", path, tmp_path / "meta.json", "abc", "2026-08-11T00:00:00Z")

    records, warnings = normalize_xiyou_xlsx(path, artifact, case)

    own_score = next(r for r in records if r.asin == case.own_asin and r.metric == "xiyou_7d_traffic_score")
    assert own_score.value == 88.5
    assert own_score.locator == "多ASIN总览!B2"
    assert own_score.source_level == "L3"
    assert warnings == []


def test_normalize_xiyou_maps_real_overview_headers(tmp_path):
    path = tmp_path / "overview.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([
        "ASIN", "价格", "Review数量", "Review评分", "7天流量得分",
        "7天自然流量得分占比", "7天广告流量得分占比",
        "流量得分（自然）", "流量得分（广告）", "词数量（自然）",
        "词数量（广告）", "近30天订单量",
    ])
    ws.append(["B0SYNTH001", 19.99, 33, 4.9, 46, 0.9565, 0.0435, 44, 2, 62, 2, 0])
    wb.save(path)
    case = CaseConfig("case-1", "US", "B0SYNTH001", (), date(2026, 7, 12), date(2026, 8, 10))
    artifact = RawArtifact("X1", "XIYOU", "L3", path, tmp_path / "meta.json", "abc", "2026-08-11T00:00:00Z")

    records, _ = normalize_xiyou_xlsx(
        path, artifact, case,
        source_period=(date(2026, 8, 3), date(2026, 8, 9)),
        period_basis="companion_keyword_export",
    )

    values = {record.metric: record for record in records}
    assert values["xiyou_price"].value == 19.99
    assert values["xiyou_review_count"].value == 33
    assert values["xiyou_organic_keyword_count"].value == 62
    assert values["xiyou_estimated_30d_orders"].value == 0
    assert values["xiyou_7d_traffic_score"].period_start == "2026-08-03"
    assert values["xiyou_estimated_30d_orders"].period_start == "2026-07-11"
    assert values["xiyou_estimated_30d_orders"].period_end == "2026-08-09"


def test_normalize_xiyou_maps_wide_keyword_export_once(tmp_path):
    path = tmp_path / "keywords.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "2026-08-03~2026-08-09——按指标分组"
    ws.append([
        "关键词 (数据来源于西柚洞察)", "翻译", "词标签", "相关性得分",
        "流量总和", "周搜索量", "B0SYNTH001-流量份额", "B0SYNTH001-流量占比",
        "B0SYNTH001-自然排名", "B0SYNTH001-SP广告排名",
        "B0SYNTH002-流量份额", "B0SYNTH002-流量占比",
        "B0SYNTH002-自然排名", "B0SYNTH002-SP广告排名",
    ])
    ws.append(["dog wipes", "宠物湿巾", "机会词", 72, 30509, 26741, 0.0016, 0.0297, 94, None, 0.3937, 0.087, 55, 9])
    duplicate_layout = wb.create_sheet("2026-08-03~2026-08-09——按ASIN分组")
    duplicate_layout.append([cell.value for cell in ws[1]])
    duplicate_layout.append([cell.value for cell in ws[2]])
    wb.save(path)
    case = CaseConfig("case-1", "US", "B0SYNTH001", ("B0SYNTH002",), date(2026, 7, 12), date(2026, 8, 10))
    artifact = RawArtifact("X2", "XIYOU", "L3", path, tmp_path / "meta.json", "abc", "2026-08-11T00:00:00Z")

    records, _ = normalize_xiyou_xlsx(path, artifact, case)

    shares = [record for record in records if record.metric == "xiyou_keyword_traffic_share"]
    assert len(shares) == 2
    own = next(record for record in shares if record.asin == case.own_asin)
    assert own.value == 0.0016
    assert own.period_start == "2026-08-03"
    assert own.period_end == "2026-08-09"
    assert own.dimensions["keyword"] == "dog wipes"
    assert own.dimensions["weekly_search_volume"] == 26741
    assert own.dimensions["organic_rank"] == 94
    assert own.locator.endswith("!G2")
    assert own.dimensions["field_locators"]["competitive_traffic_share"].endswith("!G2")
    assert own.dimensions["field_locators"]["traffic_share_of_asin"].endswith("!H2")
    assert own.dimensions["field_locators"]["organic_rank"].endswith("!I2")


def test_normalize_xiyou_maps_week_average_search_volume_header(tmp_path):
    path = tmp_path / "week-average-keywords.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "2026-07~2026-08——按指标分组"
    ws.append([
        "关键词 (数据来源于西柚洞察)", "周平均搜索量",
        "B0SYNTH001-流量份额", "B0SYNTH002-流量份额",
    ])
    ws.append(["dog wipes", 26741, 0.0016, 0.3937])
    wb.save(path)
    case = CaseConfig("case-1", "US", "B0SYNTH001", ("B0SYNTH002",), date(2026, 7, 12), date(2026, 8, 10))
    artifact = RawArtifact("X4", "XIYOU", "L3", path, tmp_path / "meta.json", "abc", "2026-08-11T00:00:00Z")

    records, _ = normalize_xiyou_xlsx(
        path, artifact, case,
        source_period=(date(2026, 7, 12), date(2026, 8, 10)),
        period_basis="xiyou_ui_1month",
    )

    own = next(record for record in records if record.asin == case.own_asin)
    assert own.dimensions["weekly_search_volume"] == 26741
    assert own.dimensions["field_locators"]["weekly_search_volume"].endswith("!B2")


def test_normalize_xiyou_handles_incorrect_worksheet_dimension(tmp_path):
    path = tmp_path / "bad-dimension.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["图片", "商品标题", "ASIN", "7天流量得分"])
    ws.append([None, "title", "B0SYNTH001", 46])
    ws.cell(row=8, column=1, value="footer")
    wb.save(path)
    patched = tmp_path / "patched.xlsx"
    with ZipFile(path, "r") as source, ZipFile(patched, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(br'<dimension ref="[^"]+"\s*/>', b'<dimension ref="A1:A1"/>', data)
            target.writestr(info, data)
    patched.replace(path)
    case = CaseConfig("case-1", "US", "B0SYNTH001", (), date(2026, 7, 12), date(2026, 8, 10))
    artifact = RawArtifact("X3", "XIYOU", "L3", path, tmp_path / "meta.json", "abc", "2026-08-11T00:00:00Z")

    records, _ = normalize_xiyou_xlsx(path, artifact, case)

    assert any(record.metric == "xiyou_7d_traffic_score" and record.value == 46 for record in records)
