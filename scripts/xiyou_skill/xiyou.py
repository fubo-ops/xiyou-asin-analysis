from __future__ import annotations

import re
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import CaseConfig, MetricRecord, RawArtifact
from .normalize import make_record


def _key(value) -> str:
    return re.sub(r"[\s_\-/（）()]+", "", str(value or "")).lower()


ASIN_ALIASES = {_key(v) for v in ("ASIN", "商品ASIN", "子ASIN")}
KEYWORD_ALIASES = {_key(v) for v in ("关键词", "keyword", "搜索词")}
FIELD_ALIASES = {
    _key("价格"): ("xiyou_price", "USD"),
    _key("Review数量"): ("xiyou_review_count", "count"),
    _key("Review评分"): ("xiyou_review_rating", "rating"),
    _key("7天流量得分"): ("xiyou_7d_traffic_score", "score"),
    _key("流量得分"): ("xiyou_7d_traffic_score", "score"),
    _key("7天自然流量得分占比"): ("xiyou_7d_organic_traffic_score_share", "percent"),
    _key("7天广告流量得分占比"): ("xiyou_7d_ad_traffic_score_share", "percent"),
    _key("流量得分（自然）"): ("xiyou_organic_traffic_score", "score"),
    _key("流量得分（广告）"): ("xiyou_ad_traffic_score", "score"),
    _key("预估30天订单"): ("xiyou_estimated_30d_orders", "estimated_count"),
    _key("近30天预估订单"): ("xiyou_estimated_30d_orders", "estimated_count"),
    _key("近30天订单量"): ("xiyou_estimated_30d_orders", "estimated_count"),
    _key("自然关键词数"): ("xiyou_organic_keyword_count", "count"),
    _key("广告关键词数"): ("xiyou_ad_keyword_count", "count"),
    _key("词数量（自然）"): ("xiyou_organic_keyword_count", "count"),
    _key("词数量（广告）"): ("xiyou_ad_keyword_count", "count"),
    _key("自然排名"): ("xiyou_organic_rank", "rank"),
    _key("广告排名"): ("xiyou_ad_rank", "rank"),
    _key("搜索量"): ("xiyou_search_volume", "estimated_count"),
    _key("流量占比"): ("xiyou_traffic_share_of_asin", "percent"),
    _key("流量份额"): ("xiyou_keyword_competitive_share", "percent"),
}


def _sheet_period(title: str) -> tuple[date, date] | None:
    match = re.search(r"(\d{4})-(\d{2})-(\d{2})~(\d{4})-(\d{2})-(\d{2})", title)
    if not match:
        return None
    values = [int(value) for value in match.groups()]
    return date(values[0], values[1], values[2]), date(values[3], values[4], values[5])


def _value(row, index_by_header: dict[str, int], header: str):
    index = index_by_header.get(header)
    return row[index].value if index is not None else None


def _normalize_keyword_sheet(worksheet, artifact, case, source_period, period_basis):
    rows = worksheet.iter_rows()
    header_row = next(rows)
    headers = [str(cell.value or "").strip() for cell in header_row]
    index_by_header = {header: index for index, header in enumerate(headers)}
    keyword_col = next((i for i, header in enumerate(headers) if header.startswith("关键词")), None)
    search_volume_header = next(
        (header for header in ("周搜索量", "周平均搜索量") if header in index_by_header), None
    )
    if keyword_col is None:
        return [], [f"{worksheet.title}: 未找到关键词表头"]
    period = source_period or _sheet_period(worksheet.title)
    period_start = (period[0] if period else case.period_start).isoformat()
    period_end = (period[1] if period else case.period_end).isoformat()
    basis = period_basis or ("workbook_sheet_name" if _sheet_period(worksheet.title) else "analysis_period_fallback")
    records = []
    for row in rows:
        keyword = str(row[keyword_col].value or "").strip()
        if not keyword:
            continue
        common = {
            "keyword": keyword,
            "translation": _value(row, index_by_header, "翻译"),
            "keyword_tag": _value(row, index_by_header, "词标签"),
            "relevance_score": _value(row, index_by_header, "相关性得分"),
            "relevance_tier": _value(row, index_by_header, "相关性档位"),
            "traffic_total": _value(row, index_by_header, "流量总和"),
            "weekly_search_volume": _value(row, index_by_header, search_volume_header) if search_volume_header else None,
            "cpc_suggested": _value(row, index_by_header, "CPC建议竞价($)"),
            "click_conversion_rate": _value(row, index_by_header, "点击转化率(均值)"),
            "period_basis": basis,
        }
        for asin in case.all_asins:
            share_header = f"{asin}-流量份额"
            share_col = index_by_header.get(share_header)
            if share_col is None:
                continue
            def cell_locator(header: str) -> str | None:
                index = index_by_header.get(header)
                if index is None:
                    return None
                return f"{worksheet.title}!{get_column_letter(index + 1)}{row[share_col].row}"

            dimensions = dict(common)
            dimensions.update({
                "traffic_share_of_asin": _value(row, index_by_header, f"{asin}-流量占比"),
                "organic_rank": _value(row, index_by_header, f"{asin}-自然排名"),
                "organic_rank_page": _value(row, index_by_header, f"{asin}-自然排名(页码-页内排名)"),
                "sp_ad_rank": _value(row, index_by_header, f"{asin}-SP广告排名"),
                "sp_ad_rank_page": _value(row, index_by_header, f"{asin}-SP广告排名(页码-页内排名)"),
                "placements": _value(row, index_by_header, f"{asin}-展示位置"),
                "metric_semantics": "keyword_competitive_traffic_share",
                "field_locators": {
                    "weekly_search_volume": (
                        f"{worksheet.title}!{row[index_by_header[search_volume_header]].coordinate}"
                        if search_volume_header else None
                    ),
                    "traffic_total": f"{worksheet.title}!{row[index_by_header['流量总和']].coordinate}" if "流量总和" in index_by_header else None,
                    "competitive_traffic_share": cell_locator(share_header),
                    "traffic_share_of_asin": cell_locator(f"{asin}-流量占比"),
                    "organic_rank": cell_locator(f"{asin}-自然排名"),
                    "organic_rank_page": cell_locator(f"{asin}-自然排名(页码-页内排名)"),
                    "sp_ad_rank": cell_locator(f"{asin}-SP广告排名"),
                    "sp_ad_rank_page": cell_locator(f"{asin}-SP广告排名(页码-页内排名)"),
                    "placements": cell_locator(f"{asin}-展示位置"),
                    "cpc_suggested": cell_locator("CPC建议竞价($)"),
                    "click_conversion_rate": cell_locator("点击转化率(均值)"),
                    "relevance_score": cell_locator("相关性得分"),
                },
            })
            records.append(make_record(
                artifact, case, asin=asin, metric="xiyou_keyword_traffic_share",
                value=row[share_col].value or 0, unit="percent",
                locator=f"{worksheet.title}!{row[share_col].coordinate}", dimensions=dimensions,
                period_start=period_start, period_end=period_end,
            ))
    return records, []


def normalize_xiyou_xlsx(
    path: str | Path, artifact: RawArtifact, case: CaseConfig, *,
    source_period: tuple[date, date] | None = None,
    period_basis: str | None = None,
) -> tuple[list[MetricRecord], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[MetricRecord] = []
    warnings: list[str] = []
    for worksheet in workbook.worksheets:
        if hasattr(worksheet, "reset_dimensions"):
            worksheet.reset_dimensions()
        if "按ASIN分组" in worksheet.title:
            continue
        first_row = next(worksheet.iter_rows(min_row=1, max_row=1), ())
        if any(str(cell.value or "").startswith("关键词") for cell in first_row):
            keyword_records, keyword_warnings = _normalize_keyword_sheet(
                worksheet, artifact, case, source_period, period_basis
            )
            records.extend(keyword_records)
            warnings.extend(keyword_warnings)
            continue
        rows = list(worksheet.iter_rows())
        header_index = next(
            (i for i, row in enumerate(rows[:10]) if any(_key(cell.value) in ASIN_ALIASES for cell in row)), None
        )
        if header_index is None:
            warnings.append(f"{worksheet.title}: 前 10 行未找到 ASIN 表头")
            continue
        headers = [_key(cell.value) for cell in rows[header_index]]
        asin_col = next(i for i, name in enumerate(headers) if name in ASIN_ALIASES)
        keyword_col = next((i for i, name in enumerate(headers) if name in KEYWORD_ALIASES), None)
        mapped_cols = {i: FIELD_ALIASES[name] for i, name in enumerate(headers) if name in FIELD_ALIASES}
        unknown = [
            rows[header_index][i].value for i, name in enumerate(headers)
            if name and name not in ASIN_ALIASES and name not in KEYWORD_ALIASES and name not in FIELD_ALIASES
        ]
        if unknown:
            warnings.append(f"{worksheet.title}: 未映射字段 {', '.join(map(str, unknown))}")
        for row in rows[header_index + 1:]:
            if len(row) <= asin_col:
                continue
            asin = str(row[asin_col].value or "").strip().upper()
            if asin not in case.all_asins:
                continue
            dimensions = {}
            if keyword_col is not None and row[keyword_col].value:
                dimensions["keyword"] = str(row[keyword_col].value).strip()
            for col, (metric, unit) in mapped_cols.items():
                if len(row) <= col:
                    continue
                value = row[col].value
                if value is None or value == "":
                    continue
                locator = f"{worksheet.title}!{row[col].coordinate}"
                metric_period = source_period
                if metric == "xiyou_estimated_30d_orders" and source_period:
                    metric_period = (source_period[1] - timedelta(days=29), source_period[1])
                dimensions["period_basis"] = period_basis or ("source_period" if source_period else "analysis_period_fallback")
                records.append(make_record(
                    artifact, case, asin=asin, metric=metric, value=value, unit=unit,
                    locator=locator, dimensions=dimensions,
                    period_start=metric_period[0].isoformat() if metric_period else None,
                    period_end=metric_period[1].isoformat() if metric_period else None,
                ))
    workbook.close()
    return records, warnings

